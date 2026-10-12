"""End-to-end tests for the Excel report and CLI."""

from __future__ import annotations

from openpyxl import load_workbook

from scroi.cli import main
from tests.conftest import DATA_DIR

EXPECTED_SHEETS = ["Inputs", "Cash Flows", "KPIs", "Sensitivity", "Optimisation Result"]


def test_evaluate_writes_report(tmp_path, capsys) -> None:
    code = main(["evaluate", str(DATA_DIR / "scenarios" / "agv_fleet.yaml"),
                 "--portfolio", str(DATA_DIR / "portfolio.yaml"), "--out", str(tmp_path)])
    assert code == 0
    xlsx = next(tmp_path.glob("*_report.xlsx"))
    wb = load_workbook(xlsx)
    assert wb.sheetnames == EXPECTED_SHEETS
    assert "SYNTHETIC" in str(wb["KPIs"]["A2"].value)
    assert len(list(tmp_path.glob("*.png"))) == 2
    out = capsys.readouterr().out
    assert "NPV" in out and "Portfolio optimisation" in out


def test_evaluate_without_portfolio(tmp_path) -> None:
    assert main(["evaluate", str(DATA_DIR / "scenarios" / "sortation_retrofit.csv"),
                 "--out", str(tmp_path)]) == 0
    wb = load_workbook(next(tmp_path.glob("*_report.xlsx")))
    assert "Not run" in str(wb["Optimisation Result"]["A4"].value)


def test_compare_and_optimise(capsys) -> None:
    assert main(["compare", str(DATA_DIR / "scenarios" / "*.yaml")]) == 0
    assert main(["optimise", str(DATA_DIR / "portfolio.yaml"), "--budget", "1000000"]) == 0
    out = capsys.readouterr().out
    assert "AGV fleet" in out and "Budget used" in out


def test_bad_input_returns_error_code(tmp_path, capsys) -> None:
    bad = tmp_path / "bad.yaml"
    bad.write_text("name: x\ncapex: -5\nproject_life_years: 3\ndiscount_rate: 0.1\nannual_labour_savings: 1\n")
    assert main(["evaluate", str(bad), "--out", str(tmp_path)]) == 2
    assert "capex" in capsys.readouterr().err


def test_formatters_handle_nan() -> None:
    from scroi.cli import _fmt_pct, _fmt_years

    assert _fmt_years(float("nan")) == "never"
    assert _fmt_pct(None) == "n/a"
