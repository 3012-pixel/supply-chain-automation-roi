"""Multi-sheet Excel report written with openpyxl.

Sheets: Inputs, Cash Flows, KPIs, Sensitivity (tornado + scenarios),
Optimisation Result.
"""

from __future__ import annotations

from dataclasses import asdict
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from scroi.config import ScenarioInputs
from scroi.metrics import KPIs
from scroi.optimisation import OptimisationResult

HEADER_FILL = PatternFill(fill_type="solid", start_color="1F3A5F", end_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, color="9C2A00")
MONEY = '#,##0;[Red]-#,##0'
PCT = '0.0%'
YEARS = '0.00'

SYNTHETIC_NOTE = "SYNTHETIC DATA - illustrative figures only, not from any real company."


def _title(ws: Worksheet, text: str, synthetic: bool) -> int:
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if synthetic:
        ws["A2"] = SYNTHETIC_NOTE
        ws["A2"].font = NOTE_FONT
    return 4  # first free row


def _write_table(ws: Worksheet, df: pd.DataFrame, start_row: int,
                 formats: dict[str, str] | None = None) -> int:
    """Write a DataFrame with a styled header; return the next free row."""
    formats = formats or {}
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=str(col))
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, (col, value) in enumerate(zip(df.columns, row), start=1):
            if isinstance(value, float) and pd.isna(value):
                value = None
            elif hasattr(value, "item"):  # numpy scalar -> python
                value = value.item()
            cell = ws.cell(row=i, column=j, value=value)
            if col in formats:
                cell.number_format = formats[col]
    return start_row + len(df) + 2


def _autosize(ws: Worksheet, min_width: int = 10, max_width: int = 45) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, w + 2))


def _inputs_frame(s: ScenarioInputs) -> pd.DataFrame:
    rows: list[tuple[str, Any]] = []
    for key, value in asdict(s).items():
        if key == "throughput":
            rows += [(f"throughput.{k}", v) for k, v in value.items()]
        elif key == "ramp_up":
            rows.append((key, ", ".join(f"{r:g}" for r in value)))
        elif key == "scenarios":
            for name, mult in value.items():
                rows.append((f"scenario.{name}", ", ".join(f"{d} x{m:g}" for d, m in mult.items()) or "as input"))
        else:
            rows.append((key, value))
    return pd.DataFrame(rows, columns=["parameter", "value"])


def write_excel_report(
    path: str | Path,
    scenario: ScenarioInputs,
    cash_flows: pd.DataFrame,
    kpis: KPIs,
    tornado_df: pd.DataFrame,
    scenarios_df: pd.DataFrame,
    optimisation: OptimisationResult | None = None,
    images: Iterable[tuple[str, str | Path]] = (),
) -> Path:
    """Write the full report workbook and return its path.

    Args:
        images: ``(sheet_name, png_path)`` pairs to embed to the right of
            that sheet's tables.
    """
    path = Path(path)
    syn = scenario.synthetic
    wb = Workbook()

    ws = wb.active
    ws.title = "Inputs"
    r = _title(ws, f"Inputs - {scenario.name}", syn)
    _write_table(ws, _inputs_frame(scenario), r)
    _autosize(ws)

    ws = wb.create_sheet("Cash Flows")
    r = _title(ws, f"Annual cash flows ({scenario.currency})", syn)
    money_cols = {c: MONEY for c in cash_flows.columns if c not in {"year", "ramp_factor", "discount_factor"}}
    _write_table(ws, cash_flows, r, {**money_cols, "ramp_factor": PCT, "discount_factor": "0.0000"})
    _autosize(ws)

    ws = wb.create_sheet("KPIs")
    r = _title(ws, "Key investment metrics", syn)
    k = kpis.as_dict()
    kpi_rows = pd.DataFrame([
        ("Net present value (NPV)", k["npv"], MONEY),
        ("Internal rate of return (IRR)", k["irr"], PCT),
        ("Simple payback (years)", k["simple_payback_years"], YEARS),
        ("Discounted payback (years)", k["discounted_payback_years"], YEARS),
        ("Profitability index", k["profitability_index"], "0.00"),
        ("Total undiscounted net cash flow", k["total_net_cash_flow"], MONEY),
        ("Discount rate", scenario.discount_rate, PCT),
        ("Recommendation", kpis.decision, "@"),
    ], columns=["metric", "value", "fmt"])
    _write_table(ws, kpi_rows[["metric", "value"]], r)
    for i, fmt in enumerate(kpi_rows["fmt"], start=r + 1):
        ws.cell(row=i, column=2).number_format = fmt
        if ws.cell(row=i, column=2).value is None:
            ws.cell(row=i, column=2, value="n/a (never reached / undefined)")
    _autosize(ws)

    ws = wb.create_sheet("Sensitivity")
    r = _title(ws, "Sensitivity analysis - NPV by driver (one at a time)", syn)
    r = _write_table(ws, tornado_df, r, {c: MONEY for c in tornado_df.columns if c.startswith(("npv", "base"))})
    ws.cell(row=r, column=1, value="Scenario analysis").font = Font(bold=True, size=12)
    _write_table(ws, scenarios_df, r + 1, {
        "npv": MONEY, "irr": PCT, "simple_payback_years": YEARS, "discounted_payback_years": YEARS,
        "profitability_index": "0.00", "total_net_cash_flow": MONEY,
    })
    _autosize(ws)

    ws = wb.create_sheet("Optimisation Result")
    r = _title(ws, "Portfolio optimisation - best mix of automation options", syn)
    if optimisation is None:
        ws.cell(row=r, column=1, value="Not run. Pass --portfolio <config.yaml> to include it.")
    else:
        o = optimisation
        summary = pd.DataFrame([
            ("Budget", o.budget, o.total_capex, o.budget - o.total_capex),
            ("Space (m2)", o.space_m2, o.total_footprint_m2, o.space_m2 - o.total_footprint_m2),
        ], columns=["constraint", "limit", "used", "slack"])
        r = _write_table(ws, summary, r, {"limit": MONEY, "used": MONEY, "slack": MONEY})
        ws.cell(row=r, column=1, value="Portfolio NPV")
        ws.cell(row=r, column=2, value=o.total_npv).number_format = MONEY
        ws.cell(row=r + 1, column=1, value="Solver status")
        ws.cell(row=r + 1, column=2, value=o.status)
        table = o.options.sort_values(["selected", "npv"], ascending=False)
        _write_table(ws, table, r + 3, {"capex": MONEY, "annual_net_benefit": MONEY, "npv": MONEY})
    _autosize(ws)

    for sheet, img_path in images:
        img = Image(str(img_path))
        img.width, img.height = img.width * 0.5, img.height * 0.5
        target = wb[sheet]
        target.add_image(img, f"{get_column_letter(target.max_column + 2)}4")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
